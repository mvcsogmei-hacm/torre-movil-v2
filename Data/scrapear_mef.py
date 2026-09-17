# -*- coding: utf-8 -*-
"""
Descarga de la Consulta Amigable del MEF (Transparencia Económica) los datos de
ejecución presupuestal que usan las pantallas INICIO, RANKING y PLIEGO MVCS de
La torre, y los escribe en las hojas correspondientes de Data.xlsx, con el mismo
formato que ya lee generar_proyectos.py. Las demás hojas del Excel no se tocan.

Uso, desde la carpeta Data/:
    py scrapear_mef.py                 (año actual)
    py scrapear_mef.py 2025            (otro año)
    py scrapear_mef.py --indicador sector
                                       (el indicador principal de Inicio con el
                                        total del SECTOR 37 en vez del PLIEGO 037)
Después:
    py generar_proyectos.py            (regenera proyectos.js)
y actualizar FECHAS_CORTE.inicio en index.html con la fecha de descarga.

Hojas que actualiza en Data.xlsx (formato = el que lee el generador):
  INDICADOR PRINCIPAL   SECTOR (% ejecución) · PIM · DEVENGADO
  PLIEGOS               ENTIDAD | PIM | DEVENGADO | EJECUCION   (MVCS, SBN, SENCICO, OTASS, COFOPRI)
  MVCS                  fila | Todo | Actividades | Proyectos   (PLIEGO, ADM GENERAL, PNSU, PNSR, PASLC)
  RANKING SECTOR        "NN: SECTOR" | PIM | % AVANCE            (todos los sectores del Gob. Nacional)
  RANKING INVERSIONES   ídem, solo proyectos  (ap=Proyecto)
  RANKING ACTIVIDADES   ídem, solo actividades (ap=Actividad)
  FUENTE MEF            año, fecha de descarga y direcciones consultadas

Son 9 consultas (3 niveles × total / proyectos / actividades), unos 30 segundos.
Adaptado de Reporte resumido/dashboard/data/scrapear_mef.py (mismo sitio, misma
forma de leer el cuadro), pero sin el desglose por departamento, que la app no usa.
Necesita openpyxl.
"""
import datetime
import html
import os
import re
import sys
import time
import urllib.request

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: py -m pip install openpyxl")

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, "Data.xlsx")
BASE = "https://apps5.mineco.gob.pe/transparencia/Navegador/Navegar_7.aspx"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")
PAUSA = 2.5          # segundos entre consultas (el sitio tiene protección contra robots)

SECTOR = "37"        # Vivienda, Construcción y Saneamiento
PLIEGO = "037"       # Ministerio de Vivienda, Construcción y Saneamiento

# Indicador principal de Inicio: "pliego" = fila 037 del nivel Pliego (así venía en
# el Excel histórico: PIM 3,865 M); "sector" = fila 37 del nivel Sector (incluye SBN,
# SENCICO, OTASS y COFOPRI; es la misma cifra con la que Vivienda aparece en el ranking).
INDICADOR_PRINCIPAL = "pliego"

# Nombre corto que muestra la app para cada pliego del sector (hoja PLIEGOS, en este orden)
PLIEGOS = [("037", "MVCS"), ("056", "SBN"), ("205", "SENCICO"), ("207", "OTASS"), ("211", "COFOPRI")]

# Unidades ejecutoras del pliego 037 que muestra el detalle "Pliego MVCS" (hoja MVCS, en este orden)
UNIDADES = [("001", "ADM GENERAL"), ("004", "PNSU"), ("005", "PNSR"), ("006", "PASLC")]

# (clave interna, valor del parámetro ap de la Consulta Amigable; None = total)
TIPOS = [("total", None), ("proyectos", "Proyecto"), ("actividades", "Actividad")]


def direccion(anio, nivel, ap=None):
    """Ruta de la Consulta Amigable dentro del Gobierno Nacional (1=E):
    nivel "sector"  -> todos los sectores;
    nivel "pliego"  -> pliegos del sector 37;
    nivel "ue"      -> unidades ejecutoras del pliego 037.
    ap=Proyecto filtra inversiones, ap=Actividad actividades, sin ap el total."""
    ruta = {"sector": "&1=E&2=",
            "pliego": "&1=E&2=" + SECTOR + "&3=",
            "ue":     "&1=E&2=" + SECTOR + "&3=" + PLIEGO + "&4="}[nivel]
    return (BASE + "?_uhc=yes&0=" + ruta + "&y=" + str(anio) + "&cpage=1&psize=400"
            + ("&ap=" + ap if ap else ""))


def descargar(url, intentos=3):
    for i in range(intentos):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "es-PE,es;q=0.9"})
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.read().decode("utf-8", "replace")
        except Exception as e:      # red caída o bloqueo temporal: reintenta con espera creciente
            if i == intentos - 1:
                raise RuntimeError("No se pudo descargar " + url + " · " + repr(e))
            time.sleep(10 * (i + 1))


def numero(texto):
    t = texto.replace(",", "").strip()
    return float(t) if t not in ("", "-") else 0.0


def leer_cuadro(pagina):
    """Filas del cuadro. Columnas del MEF: nombre, PIA, PIM, Certificación, Compromiso anual,
    Atención de compromiso mensual, Devengado, Girado, Avance %.
    Devuelve [{cod, nombre, pim, devengado, pct}] con pct en fracción (0.744)."""
    filas = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", pagina, re.S):
        if 'name="grp1"' not in tr:
            continue
        celdas = [re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", c))).strip()
                  for c in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
        celdas = [c for c in celdas if c]
        if len(celdas) < 9 or ":" not in celdas[0]:
            continue
        cod, nombre = [x.strip() for x in celdas[0].split(":", 1)]
        pim, dev, av = numero(celdas[2]), numero(celdas[6]), numero(celdas[8])
        if pim > 0 and abs(dev / pim * 100 - av) > 0.2:
            raise RuntimeError("Las columnas del cuadro cambiaron (devengado entre PIM no coincide con el "
                               "avance en '" + celdas[0] + "'). Revisar el script.")
        filas.append({"cod": cod, "nombre": nombre, "pim": round(pim), "devengado": round(dev),
                      "pct": (dev / pim) if pim > 0 else 0.0})
    return filas


def buscar(filas, cod, que):
    """Fila cuyo código empieza por `cod` (las UE vienen como '001-1082')."""
    for f in filas:
        if f["cod"] == cod or f["cod"].startswith(cod + "-"):
            return f
    sys.exit("No se encontró " + que + " (código " + cod + ") en la Consulta Amigable. No se modificó Data.xlsx.")


def rehacer_hoja(wb, nombre, filas):
    """Reemplaza el contenido de una hoja conservando su posición en el libro."""
    if nombre in wb.sheetnames:
        idx = wb.sheetnames.index(nombre)
        wb.remove(wb[nombre])
    else:
        idx = len(wb.sheetnames)
    ws = wb.create_sheet(nombre, idx)
    for fila in filas:
        ws.append(fila)
    return ws


def formatear(ws, col_pct=None, cols_monto=(), desde=2):
    for fila in ws.iter_rows(min_row=desde):
        for c in cols_monto:
            fila[c - 1].number_format = "#,##0"
        if col_pct:
            fila[col_pct - 1].number_format = "0.0%"


def guardar(anio, datos, descargado, urls):
    """datos = {"sector": {tipo: filas}, "pliego": {tipo: filas}, "ue": {tipo: filas}}"""
    print("Abriendo Data.xlsx ...")
    wb = openpyxl.load_workbook(XLSX)

    # -- INDICADOR PRINCIPAL ------------------------------------------------
    if INDICADOR_PRINCIPAL == "sector":
        p = buscar(datos["sector"]["total"], SECTOR, "el sector " + SECTOR)
    else:
        p = buscar(datos["pliego"]["total"], PLIEGO, "el pliego " + PLIEGO)
    ws = rehacer_hoja(wb, "INDICADOR PRINCIPAL", [
        [None, None, None],
        [None, "SECTOR", p["pct"]],
        [None, "PIM", p["pim"]],
        [None, "DEVENGADO", p["devengado"]],
    ])
    ws["C2"].number_format = "0.0%"
    ws["C3"].number_format = ws["C4"].number_format = "#,##0"

    # -- PLIEGOS (tarjetas "A nivel de pliegos del sector") -----------------
    filas = [[None] * 6, [None, None, "ENTIDAD", "PIM", "DEVENGADO", "EJECUCION"]]
    for cod, corto in PLIEGOS:
        f = buscar(datos["pliego"]["total"], cod, "el pliego " + corto)
        filas.append([None, None, corto, f["pim"], f["devengado"], f["pct"]])
    ws = rehacer_hoja(wb, "PLIEGOS", filas)
    formatear(ws, col_pct=6, cols_monto=(4, 5), desde=3)

    # -- MVCS (detalle Pliego MVCS: Todo / Actividades / Proyectos) ---------
    def trio(nivel, cod, que):
        return [buscar(datos[nivel][t], cod, que)["pct"] for t in ("total", "actividades", "proyectos")]
    filas = [[None] * 5, [None, None, "Todo", "Actividades", "Proyectos"],
             [None, "PLIEGO"] + trio("pliego", PLIEGO, "el pliego " + PLIEGO)]
    for cod, corto in UNIDADES:
        filas.append([None, corto] + trio("ue", cod, "la unidad ejecutora " + corto))
    ws = rehacer_hoja(wb, "MVCS", filas)
    for fila in ws.iter_rows(min_row=3):
        for c in (3, 4, 5):
            fila[c - 1].number_format = "0.0%"

    # -- RANKINGS (todos los sectores del Gobierno Nacional) ----------------
    for hoja, tipo in (("RANKING SECTOR", "total"), ("RANKING INVERSIONES", "proyectos"),
                       ("RANKING ACTIVIDADES", "actividades")):
        filas = [["Sector - Ranking por Ejecución presupuestal", "PIM", "% AVANCE"]]
        for f in datos["sector"][tipo]:
            filas.append([f["cod"] + ": " + f["nombre"], f["pim"], f["pct"]])
        ws = rehacer_hoja(wb, hoja, filas)
        formatear(ws, col_pct=3, cols_monto=(2,))
        ws.column_dimensions["A"].width = 52

    # -- FUENTE MEF (trazabilidad) ------------------------------------------
    filas = [["Dato", "Valor"],
             ["Fuente", "MEF · Consulta Amigable (Consulta de Ejecución del Gasto)"],
             ["Año", anio],
             ["Descargado", descargado],
             ["Indicador principal", "Pliego 037 (MVCS)" if INDICADOR_PRINCIPAL == "pliego"
                                     else "Sector 37 (incluye SBN, SENCICO, OTASS y COFOPRI)"],
             ["Generado por", "Data/scrapear_mef.py"]]
    for etiqueta, url in urls:
        filas.append([etiqueta, url])
    ws = rehacer_hoja(wb, "FUENTE MEF", filas)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120

    tmp = XLSX + ".tmp.xlsx"
    wb.save(tmp)
    try:
        os.replace(tmp, XLSX)
    except PermissionError:
        os.remove(tmp)
        sys.exit("No se pudo reemplazar Data.xlsx: ciérralo en Excel y vuelve a correr el script.")


def main():
    global INDICADOR_PRINCIPAL
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")   # tildes seguras en la consola de Windows
    args = sys.argv[1:]
    if "--indicador" in args:
        i = args.index("--indicador")
        INDICADOR_PRINCIPAL = args[i + 1].lower()
        del args[i:i + 2]
        if INDICADOR_PRINCIPAL not in ("pliego", "sector"):
            sys.exit("--indicador debe ser pliego o sector")
    anio = int(args[0]) if args else datetime.date.today().year
    if not os.path.exists(XLSX):
        sys.exit("No existe " + XLSX)
    print("Consulta Amigable · Gobierno Nacional · sector " + SECTOR + " / pliego " + PLIEGO + " ·", anio)

    datos, urls = {}, []
    for nivel, etiqueta in (("sector", "Sectores"), ("pliego", "Pliegos del sector 37"),
                            ("ue", "Unidades ejecutoras del pliego 037")):
        datos[nivel] = {}
        for tipo, ap in TIPOS:
            url = direccion(anio, nivel, ap)
            urls.append([etiqueta + " · " + tipo, url])
            filas = leer_cuadro(descargar(url))
            if not filas:
                sys.exit("La consulta '" + etiqueta + " · " + tipo + "' no trajo filas: la página pudo "
                         "cambiar o bloquear la consulta. No se modificó Data.xlsx.\n" + url)
            datos[nivel][tipo] = filas
            print("  " + etiqueta + " · " + tipo + ":", len(filas), "filas")
            time.sleep(PAUSA)

    if len(datos["sector"]["total"]) < 10:
        sys.exit("El cuadro de sectores trajo solo " + str(len(datos["sector"]["total"])) + " filas. No se modificó Data.xlsx.")

    # control: inversiones + actividades = total (así lo entrega el MEF)
    for nivel in datos:
        tot = {f["cod"]: f["pim"] for f in datos[nivel]["total"]}
        suma = {}
        for t in ("proyectos", "actividades"):
            for f in datos[nivel][t]:
                suma[f["cod"]] = suma.get(f["cod"], 0) + f["pim"]
        malos = [c for c in tot if tot[c] and abs(suma.get(c, 0) - tot[c]) / tot[c] > 0.005]
        if malos:
            print("  Aviso (" + nivel + "): proyectos + actividades no da el total en", ", ".join(malos))

    descargado = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    guardar(anio, datos, descargado, urls)

    v = buscar(datos["sector"]["total"], SECTOR, "el sector 37")
    p = buscar(datos["pliego"]["total"], PLIEGO, "el pliego 037")
    print("\nData.xlsx actualizado (" + descargado + ").")
    print("  Sector 37 · PIM S/ {:,.0f} · devengado S/ {:,.0f} · {:.1f}%".format(v["pim"], v["devengado"], v["pct"] * 100))
    print("  Pliego 037 · PIM S/ {:,.0f} · devengado S/ {:,.0f} · {:.1f}%".format(p["pim"], p["devengado"], p["pct"] * 100))
    orden = sorted(datos["sector"]["total"], key=lambda f: -f["pct"])
    puesto = [f["cod"] for f in orden].index(SECTOR) + 1
    print("  Ranking total: Vivienda en el puesto", puesto, "de", len(orden))
    print("\nAhora corre:  py generar_proyectos.py")
    print("y pon FECHAS_CORTE.inicio = \"" + descargado[:10] + "\" en index.html.")


if __name__ == "__main__":
    main()
